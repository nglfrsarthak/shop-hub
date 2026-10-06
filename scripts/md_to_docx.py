#!/usr/bin/env python3
"""Render docs/submission.md to a Word .docx.

A small, purpose-built Markdown subset is converted: headings, paragraphs,
bullet lists, fenced code blocks, GFM pipe tables, horizontal rules and images.

Each piece of evidence is a Field/Value table immediately followed by an image.
The renderer detects that pairing and lays the two out side by side in one
table, which keeps every screenshot to a single compact band on the page.

Usage:
    python scripts/md_to_docx.py docs/submission.md docs/ShopHub-Lab-Submission.docx
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

try:
    from PIL import Image
except Exception:  # pragma: no cover
    Image = None

MAX_IMG_W = 6.4   # inches (standalone images)
MAX_IMG_H = 1.85  # inches (tall images are scaled down to this height)
EVIDENCE_IMG_W = 3.15
CODE_FILL = "F4F4F4"
ACCENT = RGBColor(0x1F, 0x38, 0x64)
ACCENT2 = RGBColor(0x2E, 0x54, 0x96)
GREY = RGBColor(0x66, 0x66, 0x66)
RED = RGBColor(0xAA, 0x00, 0x00)

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")


# --------------------------------------------------------------- low-level bits
def set_font(run, size=None, bold=None, italic=None, color=None, name=None):
    if name:
        run.font.name = name
        rpr = run._element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rf.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def bottom_rule(paragraph, color="BBBBBB", size=6):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def add_runs(paragraph, text, size):
    """Append runs, honouring **bold**, `code` and *italic* spans."""
    for part in INLINE_RE.split(text):
        if part == "":
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = paragraph.add_run(part[2:-2].replace("`", ""))
            set_font(r, size=size, bold=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = paragraph.add_run(part[1:-1].replace("`", ""))
            set_font(r, size=size, italic=True)
        elif len(part) >= 2 and part[0] == "`" and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            set_font(r, size=size - 0.5, name="Consolas")
        else:
            r = paragraph.add_run(part)
            set_font(r, size=size)


# ---------------------------------------------------------------------- document
def new_document():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    pf = normal.paragraph_format
    pf.space_after = Pt(2)
    pf.space_before = Pt(0)
    pf.line_spacing = 0.95
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
    return doc


def add_heading(doc, level, text):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    if level == 1:
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(4)
        set_font(p.add_run(text), size=13.5, bold=True, color=ACCENT)
    elif level == 2:
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(text), size=11, bold=True, color=ACCENT2)
    else:
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(text), size=10, bold=True, color=ACCENT2)
    return p


def add_para(doc, text):
    p = doc.add_paragraph()
    add_runs(p, text, 10)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 0.95
    add_runs(p, text, 10)
    return p


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.line_spacing = 0.95
    shade(p, CODE_FILL)
    set_font(p.add_run("\n".join(lines)), name="Consolas", size=8)
    return p


def add_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    bottom_rule(p)
    return p


def fit_image(path, max_w, max_h):
    if Image is None:
        return max_w, None
    try:
        with Image.open(path) as im:
            w, h = im.size
        scale = max_w / w
        if h * scale > max_h:
            scale = max_h / h
        return w * scale, h * scale
    except Exception:
        return max_w, None


def add_table(doc, header, rows, drop_header=False):
    ncols = len(header) if header else (len(rows[0]) if rows else 1)
    total = len(rows) if drop_header else len(rows) + 1
    table = doc.add_table(rows=total, cols=ncols)
    table.style = "Table Grid"

    r = 0
    if not drop_header:
        for c, cell_text in enumerate(header):
            cell = table.cell(0, c)
            cell.text = ""
            add_runs(cell.paragraphs[0], cell_text, 9)
            for run in cell.paragraphs[0].runs:
                run.font.bold = True
        r = 1

    for row in rows:
        for c in range(ncols):
            text = row[c] if c < len(row) else ""
            cell = table.cell(r, c)
            cell.text = ""
            add_runs(cell.paragraphs[0], text, 9)
            for run in cell.paragraphs[0].runs:
                if c == 0:
                    run.font.bold = True
        r += 1

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def add_evidence(doc, fields, alt, path):
    """A screenshot on the left, its fields on the right, in one table."""
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    left = table.cell(0, 0)
    right = table.cell(0, 1)
    left.width = Inches(3.35)
    right.width = Inches(3.1)

    lc = left.paragraphs[0]
    lc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lc.paragraph_format.space_before = Pt(1)
    lc.paragraph_format.space_after = Pt(1)
    if os.path.exists(path):
        width, height = fit_image(path, EVIDENCE_IMG_W, MAX_IMG_H)
        run = lc.add_run()
        if height:
            run.add_picture(path, width=Inches(width), height=Inches(height))
        else:
            run.add_picture(path, width=Inches(width))
    else:
        set_font(lc.add_run(f"[missing image: {path}]"), italic=True, color=RED)

    first = True
    for row in fields:
        label = row[0] if len(row) > 0 else ""
        value = row[1] if len(row) > 1 else ""
        label = label.replace("**", "").replace("*", "").strip()
        p = right.paragraphs[0] if first else right.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(0.5)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 0.95
        set_font(p.add_run(label + ": "), size=8, bold=True)
        add_runs(p, value, 8)

    cap = right.add_paragraph()
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(1)
    set_font(cap.add_run(alt), size=7, italic=True, color=GREY)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)
    return table


# ----------------------------------------------------------------------- parsing
def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_sep(line):
    return bool(re.match(r"^\s*\|?\s*:?-{2,}.*$", line)) and "|" in line


def is_table_start(lines, i):
    return ("|" in lines[i]) and (i + 1 < len(lines)) and is_sep(lines[i + 1])


IMG_RE = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")


def convert(md_path, out_path):
    with open(md_path, "r", encoding="utf-8") as fh:
        text = fh.read()
    base = os.path.dirname(os.path.abspath(md_path))
    lines = text.split("\n")

    doc = new_document()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # fenced code block
        if line.strip().startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i].rstrip("\n"))
                i += 1
            i += 1
            add_code_block(doc, block)
            continue

        if line.strip() == "":
            i += 1
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            add_heading(doc, len(m.group(1)), m.group(2).strip())
            i += 1
            continue

        # horizontal rule
        if re.match(r"^\s*-{3,}\s*$", line):
            add_rule(doc)
            i += 1
            continue

        # standalone image
        m = IMG_RE.match(line)
        if m:
            alt, src = m.group(1), m.group(2)
            path = src if os.path.isabs(src) else os.path.join(base, src)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if os.path.exists(path):
                w, h = fit_image(path, MAX_IMG_W, 4.0)
                run = p.add_run()
                if h:
                    run.add_picture(path, width=Inches(w), height=Inches(h))
                else:
                    run.add_picture(path, width=Inches(w))
            else:
                set_font(p.add_run(f"[missing image: {path}]"), italic=True, color=RED)
            i += 1
            continue

        # table (possibly a side-by-side evidence band)
        if is_table_start(lines, i):
            header = split_row(lines[i])
            i += 2
            rows = []
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(split_row(lines[i]))
                i += 1
            drop = (header == [] or all(not h for h in header)
                    or [h.strip().lower() for h in header] == ["field", "value"])
            j = i
            while j < len(lines) and lines[j].strip() == "":
                j += 1
            mi = IMG_RE.match(lines[j]) if j < len(lines) else None
            if drop and mi:
                alt, src = mi.group(1), mi.group(2)
                path = src if os.path.isabs(src) else os.path.join(base, src)
                add_evidence(doc, rows, alt, path)
                i = j + 1
                continue
            add_table(doc, header, rows, drop_header=drop)
            continue

        # bullet list
        if re.match(r"^\s*[-*]\s+", line):
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                item = re.sub(r"^\s*[-*]\s+", "", lines[i]).strip()
                i += 1
                while (i < len(lines) and lines[i].strip()
                       and lines[i].startswith("  ")
                       and not re.match(r"^\s*[-*]\s+", lines[i])
                       and not lines[i].strip().startswith("```")):
                    item += " " + lines[i].strip()
                    i += 1
                add_bullet(doc, item)
            continue

        # paragraph
        buf = [line.strip()]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            if nxt.strip() == "":
                break
            if re.match(r"^(#{1,6})\s+", nxt) or nxt.strip().startswith("```"):
                break
            if (re.match(r"^\s*-{3,}\s*$", nxt) or IMG_RE.match(nxt)
                    or re.match(r"^\s*[-*]\s+", nxt)):
                break
            if is_table_start(lines, i):
                break
            buf.append(nxt.strip())
            i += 1
        add_para(doc, " ".join(buf))

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(2)
    print("wrote", convert(sys.argv[1], sys.argv[2]))